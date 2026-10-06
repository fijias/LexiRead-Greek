"""Optional user-maintained known-word lists for card selection."""

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.cards.formatter import normalize_card_form

# Columns of the study list that let the user mark words as known.
KNOWN_HEADERS = {"Знаю", "Known"}
LEMMA_HEADERS = {"Лемма", "Lemma"}
WORD_FILE_SUFFIXES = (".xlsx", ".txt", ".csv")


def read_known_words(path: Path) -> set[str]:
    if not path.is_file():
        raise ValueError(f"Known dictionary not found: {path}")
    return {normalize_card_form(word) for word in read_word_file(path)}


def read_word_file(path: Path) -> list[str]:
    """Words from the first column of an XLSX/CSV file or one per line of a TXT file."""
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = [row[0] for row in workbook.active.iter_rows(min_col=1, max_col=1, values_only=True) if row]
        finally:
            workbook.close()
    else:
        from src.preprocessing.loader import read_text_auto

        text = read_text_auto(path)[0]
        if path.suffix.lower() == ".csv":
            values = [row[0] for row in csv.reader(text.splitlines()) if row]
        else:
            values = text.splitlines()
    return [value.strip() for value in values if isinstance(value, str) and value.strip()]


def expand_with_lemmas(words: set[str], lexicon=None) -> set[str]:
    """Known words may be written in any form (είπε, δρόμου): add their dictionary lemmas."""
    if lexicon is None:
        return set(words)
    return set(words) | {lemma for word in words for lemma in lexicon.lemmas_of(word)}


def load_known_words(path: Path, lexicon_path=None) -> set[str]:
    words = read_known_words(path)
    if not lexicon_path or not Path(lexicon_path).is_file():
        return words
    from src.languages.greek_lexicon import GreekLexicon

    lexicon = GreekLexicon(Path(lexicon_path))
    try:
        return expand_with_lemmas(words, lexicon)
    finally:
        lexicon.close()


def append_words(dictionary: Path, words) -> int:
    """Add new words to the first column of the known-words workbook; return how many were added."""
    if dictionary.is_file():
        workbook = load_workbook(dictionary)
    else:
        workbook = Workbook()
        workbook.active.title = "Known words"
        workbook.active.append(["Known words"])
    try:
        sheet = workbook.active
        present = {
            normalize_card_form(value)
            for (value,) in sheet.iter_rows(min_col=1, max_col=1, values_only=True)
            if isinstance(value, str)
        }
        added = 0
        for word in words:
            key = normalize_card_form(word)
            if key and key not in present:
                sheet.append([word.strip()])
                present.add(key)
                added += 1
        if added:
            dictionary.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(dictionary)
        return added
    finally:
        workbook.close()


def import_word_file(source: Path, dictionary: Path) -> int:
    return append_words(dictionary, read_word_file(source))


def marked_words(study_list: Path) -> list[str]:
    """Lemmas the user marked in the "Known" column of a previous study list."""
    if not study_list.is_file():
        return []
    workbook = load_workbook(study_list, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = list(next(rows, []))
        known = next((i for i, h in enumerate(headers) if h in KNOWN_HEADERS), None)
        lemma = next((i for i, h in enumerate(headers) if h in LEMMA_HEADERS), None)
        if known is None or lemma is None:
            return []
        return [
            row[lemma]
            for row in rows
            if len(row) > max(known, lemma) and str(row[known] or "").strip() and isinstance(row[lemma], str)
        ]
    finally:
        workbook.close()


def import_marked_words(study_list: Path, dictionary: Path) -> int:
    words = marked_words(study_list)
    return append_words(dictionary, words) if words else 0
