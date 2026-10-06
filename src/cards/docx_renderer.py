import copy
import logging
import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from src.cards.formatter import (
    card_content,
    english_back_forms,
    filter_observed_forms_for_back,
    format_observed_forms,
    get_card_front_text,
    get_forms_font_size,
    get_translation_font_size,
)
from src.cards.selector import select_cards
from src.storage import atomic_path, file_hash, read_rows

logger = logging.getLogger(__name__)


def table_order(sheet_count: int) -> list[tuple[str, int]]:
    return [("front", i) for i in range(sheet_count)] + [("back", i) for i in reversed(range(sheet_count))]


def back_position(row: int, column: int, rows: int) -> tuple[int, int]:
    return rows - 1 - row, column


def _trailing_paragraph():
    """Return Word's required post-table paragraph without a visible extra line."""
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}before", "0")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}after", "0")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line", "1")
    spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule", "exact")
    properties.append(spacing)
    paragraph.append(properties)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "2")
    run_properties.append(size)
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text.text = " "
    run.append(text)
    paragraph.append(run)
    return paragraph


def _clone_template_tables(document, sheet_count: int):
    if len(document.tables) < 2:
        raise ValueError("Cards template must contain front and back tables.")
    front, back = document.tables[:2]
    if not front.rows or not back.rows:
        raise ValueError("Cards template tables must not be empty.")
    front_xml, back_xml = copy.deepcopy(front._tbl), copy.deepcopy(back._tbl)
    body = document._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)
    # A section-properties element must be the final body child.  Appending cloned
    # tables after it makes Word repair the document and can create blank pages.
    body.remove(sect_pr)
    for side, _ in table_order(sheet_count):
        body.append(copy.deepcopy(front_xml if side == "front" else back_xml))
    body.append(_trailing_paragraph())
    body.append(sect_pr)
    return document.tables


def _clear_cell(cell):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return paragraph


def _write_paragraph(paragraph, text: str, size: int, space_after: int = 3):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run(text)
    run.font.size = Pt(size)


def render(cards, template: Path, target: Path, config, language="es", machine_translation=False) -> int:
    if not template.is_file():
        raise ValueError(f"Cards template not found: {template}")
    if config.cards_per_page != config.rows * config.columns:
        raise ValueError("cards_per_page must equal rows * columns")
    sheets = math.ceil(len(cards) / config.cards_per_page) if cards else 0
    document = Document(template)
    from src.export.excel_exporter import translation_credit
    from src.output_terms import out

    document.core_properties.comments = "" if machine_translation else translation_credit()
    if sheets == 0:
        for element in list(document.element.body):
            if element.tag != qn("w:sectPr"):
                document.element.body.remove(element)
        document.add_paragraph(out("Нет слов для изучения при выбранных настройках."))
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
        return 0
    tables = _clone_template_tables(document, sheets)
    if any(len(table.rows) != config.rows or len(table.columns) != config.columns for table in tables):
        raise ValueError(f"Cards template tables must be {config.rows} × {config.columns}.")
    lemma_fallbacks = []
    for index, entry in enumerate(cards):
        sheet, local = divmod(index, config.cards_per_page)
        row, column = divmod(local, config.columns)
        front = tables[sheet].cell(row, column)
        front_text, ipa_text, translations, forms_text = card_content(entry, config.max_forms, language)
        if front_text == entry.lemma:
            lemma_fallbacks.append(f"{entry.lemma}/{entry.pos}")
        _write_paragraph(_clear_cell(front), front_text, config.front_font_size, 0)

        back_table = tables[len(tables) - sheet - 1]
        back_row, back_column = back_position(row, column, config.rows)
        back = back_table.cell(back_row, back_column)
        paragraph = _clear_cell(back)
        blocks = [(ipa_text, config.back_font_size)]
        blocks.extend(
            (
                translation,
                get_translation_font_size(
                    translation,
                    config.back_font_size,
                    config.long_translation_threshold,
                    config.long_translation_font_size,
                ),
            )
            for translation in translations
        )
        if forms_text:
            blocks.append(
                (
                    forms_text,
                    get_forms_font_size(
                        forms_text,
                        config.back_font_size,
                        config.long_forms_threshold,
                        config.long_forms_font_size,
                    ),
                )
            )
        for block_index, (text, size) in enumerate(blocks):
            _write_paragraph(paragraph if block_index == 0 else back.add_paragraph(), text, size)
    target.parent.mkdir(parents=True, exist_ok=True)
    with atomic_path(target) as temporary:
        document.save(temporary)
    if lemma_fallbacks:
        examples = ", ".join(lemma_fallbacks[:5])
        suffix = f"; примеры: {examples}" if examples else ""
        logger.info("Лемма использована на лицевой стороне для %d карточек%s.", len(lemma_fallbacks), suffix)
    return sheets


def run(
    paths,
    template,
    target,
    cards_config,
    translation_config,
    language="es",
    known_dictionary=None,
    machine_translation=False,
):
    data = {name: read_rows(path) for name, path in paths.items()}
    known_words = set()
    if known_dictionary:
        from src.cards.known_words import read_known_words

        known_words = read_known_words(Path(known_dictionary))
    cards, selection = select_cards(
        data,
        translation_config.cumulative_coverage_limit,
        translation_config.specificity_threshold,
        translation_config.min_book_occurrences,
        language,
        known_words,
    )
    sheets = render(cards, template, target, cards_config, language, machine_translation)
    from src.cards.learning_list import list_filename, render_learning_list

    list_target = target.parent / list_filename(language)
    render_learning_list(cards, list_target, cards_config, language, machine_translation)
    logger.info("Список для изучения: %s", list_target)
    long_forms = sum(
        len(
            format_observed_forms(
                english_back_forms(card, cards_config.max_forms)
                if language == "en"
                else filter_observed_forms_for_back(
                    card.observed_forms, get_card_front_text(card), cards_config.max_forms, True
                )
            )
        )
        > cards_config.long_forms_threshold
        for card in cards
    )
    logger.info(
        "Generating printable cards: requested %.2f%%, actual %.2f%%; %d cards, %d sheets, %d pages; %d forms blocks at %d pt; %s",
        selection.requested_coverage,
        selection.actual_coverage * 100,
        len(cards),
        sheets,
        sheets * 2,
        long_forms,
        cards_config.long_forms_font_size,
        target,
    )
    return {
        "selected_cards": len(cards),
        "requested_coverage": selection.requested_coverage,
        "actual_coverage": selection.actual_coverage * 100,
        "sheet_count": sheets,
        "page_count": sheets * 2,
        "template": str(template),
        "template_sha256": file_hash(template),
    }
