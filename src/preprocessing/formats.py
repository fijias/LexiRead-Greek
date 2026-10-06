"""Plain text from the supported book and subtitle formats.

Every reader returns text with one paragraph (or subtitle cue) per line, so the
normalizer and the chunker treat all formats like a TXT file.
"""

import html
import posixpath
import re
import zipfile
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

SUBTITLES = {".srt", ".vtt", ".ass", ".ssa"}
SUPPORTED_SUFFIXES = (".txt", ".epub", ".fb2", ".docx", ".pdf", *sorted(SUBTITLES))
# Scripts expected in a text of each language, to recognize PDFs with broken font encodings.
LANGUAGE_LETTERS = {"el": re.compile(r"[\u0370-\u03ff\u1f00-\u1fff]")}
_BLOCK_TAGS = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "blockquote", "section", "title"}


class SourceFormatError(ValueError):
    """A readable message for the user about a file that cannot be used."""


def suffix(path: Path) -> str:
    name = path.name.lower()
    return ".fb2" if name.endswith(".fb2.zip") else path.suffix.lower()


def read_source(source: Path, language: str, decode) -> tuple[str, str]:
    """Return (text, description); ``decode`` reads a TXT file with encoding detection."""
    kind = suffix(source)
    if kind == ".txt":
        return decode(source, language)
    readers = {".epub": read_epub, ".fb2": read_fb2, ".docx": read_docx, ".pdf": read_pdf}
    if kind in SUBTITLES:
        text, encoding = decode(source, language)
        return read_subtitles(text, kind), f"{kind[1:]} ({encoding})"
    if kind not in readers:
        raise SourceFormatError(
            f"Формат {source.suffix or source.name} не поддерживается. Поддерживаются: "
            + ", ".join(s[1:].upper() for s in SUPPORTED_SUFFIXES)
        )
    try:
        text = readers[kind](source, language) if kind == ".pdf" else readers[kind](source)
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError) as exc:
        raise SourceFormatError(f"Файл {source.name} повреждён или не является {kind[1:].upper()}.") from exc
    if not text.strip():
        raise SourceFormatError(f"В файле {source.name} не найден текст.")
    return text, kind[1:]


class _TextCollector(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "head"}:
            self.skip += 1
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "head"}:
            self.skip = max(0, self.skip - 1)
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_text(markup: str) -> str:
    collector = _TextCollector()
    collector.feed(markup)
    lines = (" ".join(line.split()) for line in "".join(collector.parts).split("\n"))
    return "\n".join(line for line in lines if line)


def read_epub(path: Path) -> str:
    with zipfile.ZipFile(path) as book:
        container = ElementTree.fromstring(book.read("META-INF/container.xml"))
        opf_path = next(e.get("full-path") for e in container.iter() if e.tag.endswith("rootfile"))
        opf = ElementTree.fromstring(book.read(opf_path))
        base = posixpath.dirname(opf_path)
        items = {
            e.get("id"): e.get("href")
            for e in opf.iter()
            if e.tag.endswith("}item") and "html" in (e.get("media-type") or "")
        }
        spine = [e.get("idref") for e in opf.iter() if e.tag.endswith("}itemref")]
        chapters = []
        for idref in spine:
            if idref in items:
                name = posixpath.normpath(posixpath.join(base, html.unescape(items[idref]).split("#")[0]))
                chapters.append(html_text(book.read(name).decode("utf-8", errors="replace")))
    return "\n\n".join(chapter for chapter in chapters if chapter)


def read_fb2(path: Path) -> str:
    data = path.read_bytes()
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            name = next(n for n in archive.namelist() if n.lower().endswith(".fb2"))
            data = archive.read(name)
    root = ElementTree.fromstring(data)
    paragraphs = []
    for body in root:
        # Footnote bodies interrupt the text with commentary; the main body comes first.
        if not body.tag.endswith("}body") or body.get("name") in {"notes", "comments"}:
            continue
        for element in body.iter():
            if element.tag.rsplit("}", 1)[-1] in {"p", "v", "subtitle", "text-author"}:
                text = " ".join("".join(element.itertext()).split())
                if text:
                    paragraphs.append(text)
    return "\n".join(paragraphs)


def read_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(p.strip() for p in paragraphs if p.strip())


_SRT_INDEX = re.compile(r"^\d+$")
_TIMECODE = re.compile(r"^\d{1,2}:\d{2}(?::\d{2})?[.,]\d{1,3}\s*-->\s*\d{1,2}:\d{2}")
_MARKUP = re.compile(r"<[^>]+>|\{[^}]*\}")


def read_subtitles(text: str, kind: str) -> str:
    """One line per cue: indexes, timecodes, styling tags and speaker dashes removed."""
    if kind in {".ass", ".ssa"}:
        cues = []
        for line in text.splitlines():
            if line.startswith("Dialogue:"):
                # Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text
                cues.append(line.split(",", 9)[-1].replace("\\N", " ").replace("\\n", " "))
    else:
        cues, current = [], []
        for line in text.splitlines() + [""]:
            line = line.strip()
            if not line:
                if current:
                    cues.append(" ".join(current))
                current = []
            elif line == "WEBVTT" or line.startswith(("NOTE", "STYLE")) or _TIMECODE.match(line):
                continue
            elif _SRT_INDEX.match(line) and not current:
                continue
            else:
                current.append(line)
    cleaned = (" ".join(_MARKUP.sub("", cue).replace("&nbsp;", " ").split()).lstrip("-– ") for cue in cues)
    return "\n".join(cue for cue in cleaned if cue)


_HYPHENATED = re.compile(r"(\w)[-\u00ad]\n(\w)")
# Fewer characters per page on average means the PDF is a scan without a text layer.
MIN_CHARS_PER_PAGE = 20
_PAGE_NUMBER = re.compile(r"^(?:[-–— ]*\d+[-–— ]*|\d+\s*/\s*\d+)$")


def read_pdf(path: Path, language: str = "el") -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise SourceFormatError(f"PDF {path.name} защищён паролем.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise SourceFormatError(f"Не удалось прочитать PDF {path.name}: файл повреждён.") from exc
    if sum(len(page.strip()) for page in pages) < MIN_CHARS_PER_PAGE * max(1, len(pages)):
        raise SourceFormatError(
            f"В PDF {path.name} почти нет текстового слоя — вероятно, это скан. "
            "Распознайте текст (OCR) и сохраните книгу в TXT, EPUB или PDF с текстом."
        )
    text = join_pdf_pages(pages)
    letters = re.findall(r"[^\W\d_]", text)
    expected = LANGUAGE_LETTERS.get(language)
    if expected and letters and sum(bool(expected.match(c)) for c in letters) < 0.3 * len(letters):
        raise SourceFormatError(
            f"Текст PDF {path.name} извлекается неверными символами (шрифты без таблицы Unicode). "
            "Откройте файл в Calibre или Word и сохраните как TXT или EPUB."
        )
    return text


def join_pdf_pages(pages: list[str]) -> str:
    """Drop running headers, footers and page numbers; rejoin hyphenated and wrapped lines."""
    page_lines = [[line.strip() for line in page.splitlines() if line.strip()] for page in pages]
    # A line repeated at the top or bottom of many pages is a running header or footer.
    edges = Counter(
        line for lines in page_lines for line in {*lines[:2], *lines[-2:]} if not _PAGE_NUMBER.match(line)
    )
    repeated = {line for line, count in edges.items() if len(pages) >= 4 and count >= len(pages) / 2}
    kept = []
    for lines in page_lines:
        kept.extend(line for line in lines if line not in repeated and not _PAGE_NUMBER.match(line))
    text = _HYPHENATED.sub(r"\1\2", "\n".join(kept))
    # PDF lines are visual; a paragraph ends where a line ends with sentence punctuation.
    return re.sub(r"(?<![.!;:?»…\"')])\n(?=\S)", " ", text)
