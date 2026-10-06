import csv
import logging
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.export.review import REVIEW_COLUMNS, make_review_rows
from src.grammar.pos_mapping import POS_RU
from src.lemma_groups import group_lemmas
from src.storage import atomic_path, read_rows

logger = logging.getLogger(__name__)

KAIKKI_TRANSLATION_CREDIT = "Переводы: Kaikki.org / Wiktionary contributors, CC BY-SA 4.0. "
WORDFREQ_CREDIT = "Частотность: wordfreq, Copyright 2022 Robyn Speer."

LEMMA_COLUMNS = [
    "Ранг",
    "Лемма",
    "Учебная форма",
    "Суммарное число вхождений",
    "Всего вхождений леммы",
    "Кумулятивное покрытие",
    "Относительная частота в книге",
    "Частота Wordfreq",
    "Специфичность",
    "Количество контекстных блоков",
    "Транскрипция",
    "Часть речи",
    "Род",
    "Грамматические формы",
    "Инфинитив",
    "Группа спряжения",
    "Регулярность",
    "Словоформы в книге",
    "Перевод на русский",
    "Перевод на английский",
    "Пример из книги",
]
FORM_COLUMNS = [
    "Ранг",
    "Словоформа",
    "Число вхождений",
    "Относительная частота в книге",
    "Частота Wordfreq",
    "Транскрипция",
    "Лемма",
    "Часть речи",
    "Грамматические признаки",
    "Перевод на русский",
    "Перевод на английский",
    "Пример из книги",
]


def make_tables(data):
    lg = {r["id"]: r for r in data["lemma_grammar"]}
    fg = {r["id"]: r for r in data["form_grammar"]}
    ipa = {r["text"]: f"/{r['ipa']}/" if r["ipa"] else "" for r in data["ipa"]}
    tr = {r["id"]: r for r in data["translations"]}
    lemma_rows, form_rows = [], []
    groups = group_lemmas(data["lemmas"], data["forms"])
    group_by_lemma = {g["lemma"]: g for g in groups}
    for r in [row for group in groups for row in group["rows"]]:
        group = group_by_lemma[r["lemma"]]
        g, t = lg[r["id"]], tr[r["id"]]
        observed = sorted(r["observed_forms"].items(), key=lambda x: (-x[1], x[0]))
        lemma_rows.append(
            [
                group["rank"],
                r["lemma"],
                g["learning_form"],
                r["count"],
                group["count"],
                group["cumulative_coverage"],
                group["share"],
                group["reference_frequency"],
                group["specificity"],
                r["chunk_count"],
                ipa.get(r["lemma"], ""),
                POS_RU.get(r["pos"], r["pos"]),
                g["grammatical_gender"],
                g["grammar_forms"],
                g["infinitive"],
                g["conjugation_group"],
                g["regularity"],
                ", ".join(f"{f} ({n})" for f, n in observed),
                t["ru"],
                t["en"],
                r["contexts"][0] if r["contexts"] else "",
            ]
        )
    for r in sorted(data["forms"], key=lambda r: (-r["count"], r["form"], r["lemma"], r["pos"])):
        g, t = fg[r["id"]], tr[r["id"]]
        form_rows.append(
            [
                r["rank"],
                r["form"],
                r["count"],
                r.get("book_relative_frequency", 0.0),
                r.get("reference_frequency", 0.0),
                ipa.get(r["form"], ""),
                r["lemma"],
                POS_RU.get(r["pos"], r["pos"]),
                g["morph_description"],
                t["ru"],
                t["en"],
                r["contexts"][0] if r["contexts"] else "",
            ]
        )
    review_rows = make_review_rows(data)
    return [
        ("Леммы", LEMMA_COLUMNS, lemma_rows),
        ("Словоформы", FORM_COLUMNS, form_rows),
        ("Проверка", REVIEW_COLUMNS, review_rows),
    ]


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if len(value) > 32767:
            logger.warning("Excel: текст длиннее лимита ячейки; полный текст остаётся в Parquet")
            value = value[:32764] + "..."
    return value


EN_OPTIONAL_COLUMNS = {"Род", "Группа спряжения", "Регулярность", "Перевод на английский"}
EN_EXCLUDED_COLUMNS = {"Учебная форма", "Инфинитив"}
EN_WIDTHS_PX = {
    "Ранг": 40,
    **dict.fromkeys(["Лемма", "Транскрипция", "Часть речи"], 180),
    **dict.fromkeys(
        [
            "Суммарное число вхождений",
            "Всего вхождений леммы",
            "Кумулятивное покрытие",
            "Относительная частота в книге",
            "Частота Wordfreq",
            "Специфичность",
            "Количество контекстных блоков",
        ],
        120,
    ),
}


def english_columns(headers, rows):
    """Drop duplicate verb forms and optional empty display columns."""
    keep = [
        i
        for i, header in enumerate(headers)
        if header not in EN_EXCLUDED_COLUMNS
        and (header not in EN_OPTIONAL_COLUMNS or any(str(clean(row[i])).strip() for row in rows))
    ]
    return [headers[i] for i in keep], [[row[i] for i in keep] for row in rows]


def write_excel(target, tables, config, language=None, attribution=None, description=None):
    book = Workbook()
    if description is not None:
        # Whitespace makes openpyxl add xml:space, which OPC core properties reject.
        book.properties.description = description.strip()
    book.remove(book.active)
    for name, headers, rows in tables:
        if len(rows) > 1048575:
            raise ValueError(f"{name}: превышен лимит строк Excel")
        if language == "en" and name in {"Леммы", "Словоформы"}:
            headers, rows = english_columns(headers, rows)
        sheet = book.create_sheet(name)
        sheet.append(headers)
        widths = []
        for col, header in enumerate(headers, 1):
            width = (
                config.context_width
                if "Пример" in header
                else 48
                if header
                in {
                    "Словоформы в книге",
                    "Грамматические формы",
                    "Грамматические признаки",
                    "Описание",
                    "Что проверить",
                    "Что сделать",
                }
                else 26
            )
            if language == "en" and name in {"Леммы", "Словоформы"} and header in EN_WIDTHS_PX:
                # OOXML width already includes padding (Calibri 11, 7px digit at 96 DPI).
                width = math.floor(EN_WIDTHS_PX[header] / 7 * 256) / 256
            widths.append(width)
            sheet.column_dimensions[get_column_letter(col)].width = width
        for values in rows:
            sheet.append([clean(v) for v in values])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.zoomScale = 85
        for row in sheet:
            line_count = 1
            for cell in row:
                if isinstance(cell.value, str):
                    # Book text and API output must never become executable Excel formulas.
                    cell.data_type = "s"
                    line_count = max(
                        line_count, math.ceil(len(cell.value) / max(8, widths[cell.column - 1] - 3))
                    )
                cell.font = Font(name="Calibri", size=config.font_size, bold=cell.row == 1)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="right" if isinstance(cell.value, (int, float)) else "left",
                )
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="E8EDF2")
                elif headers[cell.column - 1] in {
                    "Кумулятивное покрытие",
                    "Относительная частота в книге",
                }:
                    cell.number_format = "0.00%"
                elif headers[cell.column - 1] == "Частота Wordfreq":
                    cell.number_format = "0.000000%"
                elif headers[cell.column - 1] == "Специфичность":
                    cell.number_format = "0.000000"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
            sheet.row_dimensions[row[0].row].height = min(
                409, max(32 if row[0].row == 1 else 30, line_count * 15)
            )
    if attribution:
        sheet = book.create_sheet("Атрибуция")
        sheet.column_dimensions["A"].width = 110
        for row, line in enumerate(attribution.splitlines(), 1):
            cell = sheet.cell(row=row, column=1, value=line)
            cell.font = Font(name="Calibri", size=config.font_size)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    with atomic_path(target) as tmp:
        book.save(tmp)
    book.close()


def make_attribution(source_name, machine_translation=False):
    translation = (
        "Переводы: LLM"
        if machine_translation
        else """Переводы: Kaikki.org / участники Wiktionary.
Данные автоматически отобраны и обработаны программой LexiRead Greek.
Лицензия: CC BY-SA 4.0.
https://kaikki.org/
https://en.wiktionary.org/wiki/Wiktionary:Copyrights
https://creativecommons.org/licenses/by-sa/4.0/"""
    )
    return f"""Создано с помощью LexiRead Greek — изменённой версии WordByHeart, GPL-3.0-only.
https://github.com/fijias/LexiRead-Greek
WordByHeart — Copyright © 2026 Egor Tatarnikov, GPL-3.0-only.
https://github.com/EgorTatarnikov/WordByHeart

{translation}

Общеязыковая частотность: wordfreq.
Copyright © 2022 Robyn Speer.
https://github.com/rspeer/wordfreq
https://github.com/rspeer/wordfreq/blob/master/NOTICE.md

Транскрипция IPA: eSpeak NG.
https://github.com/espeak-ng/espeak-ng

NLP-анализ: spaCy.
https://spacy.io/

Исходный текст: {source_name}.
Права на исходный текст принадлежат его автору или правообладателю.
""".lstrip()


def run(paths, output, config, language="es", source_name="", machine_translation=False):
    tables = make_tables({key: read_rows(path) for key, path in paths.items()})
    output.mkdir(parents=True, exist_ok=True)
    attribution = make_attribution(source_name, machine_translation)
    description = WORDFREQ_CREDIT
    if not machine_translation:
        description = KAIKKI_TRANSLATION_CREDIT + description
    with atomic_path(output / "ATTRIBUTION.txt") as tmp:
        tmp.write_text(attribution, encoding="utf-8")
    if config.xlsx:
        from src.languages import get_profile

        write_excel(
            output / get_profile(language).export_filename,
            tables,
            config,
            language,
            attribution,
            description,
        )
    if config.csv:
        for (_, headers, rows), name in zip(tables[:2], ("lemmas.csv", "forms.csv")):
            with atomic_path(output / name) as tmp, tmp.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                for row in rows:
                    writer.writerow(
                        [
                            "'" + v
                            if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@"))
                            else v
                            for v in row
                        ]
                    )
    logger.info("Экспорт: %s", output)
