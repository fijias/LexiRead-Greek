import zipfile

import pytest

from src.config import Preprocess
from src.preprocessing import formats
from src.preprocessing.formats import SourceFormatError, join_pdf_pages, read_subtitles
from src.preprocessing.loader import read_text_auto, run

OPF = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
  <manifest>
    <item id="c2" href="text/two.xhtml" media-type="application/xhtml+xml"/>
    <item id="c1" href="text/one.xhtml" media-type="application/xhtml+xml"/>
    <item id="css" href="style.css" media-type="text/css"/>
  </manifest>
  <spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>"""
CONTAINER = """<?xml version="1.0"?>
<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>"""


def chapter(body):
    return f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>x</title><style>p{{}}</style></head><body>{body}</body></html>'


def test_epub_follows_spine_order(tmp_path):
    path = tmp_path / "book.epub"
    with zipfile.ZipFile(path, "w") as book:
        book.writestr("mimetype", "application/epub+zip")
        book.writestr("META-INF/container.xml", CONTAINER)
        book.writestr("OEBPS/content.opf", OPF)
        book.writestr("OEBPS/text/one.xhtml", chapter("<h1>Κεφάλαιο Α'</h1><p>Ο Αντώνης ήταν <i>πολύ</i> σκάνταλος.</p>"))
        book.writestr("OEBPS/text/two.xhtml", chapter("<p>Δεύτερο&nbsp;κεφάλαιο.</p>"))
    text, kind = formats.read_source(path, "el", read_text_auto)
    assert kind == "epub"
    assert text == "Κεφάλαιο Α'\nΟ Αντώνης ήταν πολύ σκάνταλος.\n\nΔεύτερο κεφάλαιο."


def test_fb2_skips_notes_and_binary(tmp_path):
    path = tmp_path / "book.fb2"
    path.write_bytes(
        """<?xml version="1.0" encoding="windows-1253"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">
  <description><title-info><book-title>Τίτλος</book-title></title-info></description>
  <body><section><title><p>Κεφάλαιο</p></title><p>Καλημέρα <emphasis>κόσμε</emphasis>.</p></section></body>
  <body name="notes"><section><p>Σημείωση</p></section></body>
  <binary id="cover">AAAA</binary>
</FictionBook>""".encode("cp1253")
    )
    assert formats.read_source(path, "el", read_text_auto)[0] == "Κεφάλαιο\nΚαλημέρα κόσμε."


def test_zipped_fb2(tmp_path):
    path = tmp_path / "book.fb2.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("book.fb2", '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0"><body><p>Γεια</p></body></FictionBook>')
    assert formats.read_source(path, "el", read_text_auto)[0] == "Γεια"


def test_docx_paragraphs_and_tables(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Πρώτη παράγραφος.")
    document.add_paragraph("")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Κελί"
    path = tmp_path / "book.docx"
    document.save(path)
    assert formats.read_source(path, "el", read_text_auto)[0] == "Πρώτη παράγραφος.\nΚελί"


def test_srt_removes_indexes_timecodes_and_tags():
    srt = "1\n00:00:01,000 --> 00:00:03,500\n<i>Γεια σου,</i>\nΜαρία!\n\n2\n00:00:04,000 --> 00:00:05,000\n- 12 χρόνια;\n"
    assert read_subtitles(srt, ".srt") == "Γεια σου, Μαρία!\n12 χρόνια;"


def test_vtt_and_ass():
    vtt = "WEBVTT\n\nNOTE comment\n\n00:01.000 --> 00:02.000 align:start\n<v Νίκος>Καλησπέρα\n"
    assert read_subtitles(vtt, ".vtt") == "Καλησπέρα"
    ass = "[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\n" \
          "Dialogue: 0,0:00:01.00,0:00:02.00,Default,,0,0,0,,{\\i1}Τι κάνεις,{\\i0}\\Nφίλε;\n"
    assert read_subtitles(ass, ".ass") == "Τι κάνεις, φίλε;"


def test_subtitles_go_through_the_loader(tmp_path):
    source, target = tmp_path / "episode.srt", tmp_path / "out.txt"
    source.write_bytes("1\n00:00:01,000 --> 00:00:02,000\nΚαλημέρα σ' όλους\n".encode("cp1253"))
    run(source, target, Preprocess(), "el")
    assert target.read_text(encoding="utf-8") == "Καλημέρα σε όλους\n"


def test_pdf_pages_drop_headers_numbers_and_hyphenation():
    pages = [
        f"Τρελαντώνης\nΣελίδα {n}: ο Αντώνης ήταν πολύ σκάν-\nταλος και άτακτος\nκαι έτρεχε.\nΗ θεία φώναξε {n} φορές.\n- {n} -"
        for n in range(1, 6)
    ]
    lines = join_pdf_pages(pages).splitlines()
    assert lines[:2] == ["Σελίδα 1: ο Αντώνης ήταν πολύ σκάνταλος και άτακτος και έτρεχε.", "Η θεία φώναξε 1 φορές."]
    assert len(lines) == 10 and not any("Τρελαντώνης" in line or line.startswith("-") for line in lines)


class FakePage:
    def __init__(self, text):
        self.text = text

    def extract_text(self):
        return self.text


class FakeReader:
    pages = []
    is_encrypted = False

    def __init__(self, path):
        pass


@pytest.mark.parametrize(
    ("pages", "message"),
    [(["", " "], "скан"), (["Ãåéá óïõ êüóìå, ôé êÜíåéò óÞìåñá;"] * 3, "неверными символами")],
)
def test_unusable_pdfs_are_explained(tmp_path, monkeypatch, pages, message):
    import pypdf

    monkeypatch.setattr(FakeReader, "pages", [FakePage(p) for p in pages])
    monkeypatch.setattr(pypdf, "PdfReader", FakeReader)
    path = tmp_path / "book.pdf"
    path.write_bytes(b"%PDF-1.4")
    with pytest.raises(SourceFormatError, match=message):
        formats.read_source(path, "el", read_text_auto)


def test_unsupported_format(tmp_path):
    path = tmp_path / "book.mobi"
    path.write_bytes(b"x")
    with pytest.raises(SourceFormatError, match="не поддерживается"):
        formats.read_source(path, "el", read_text_auto)
