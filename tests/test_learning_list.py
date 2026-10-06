from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook

from src.cards.docx_renderer import render
from src.cards.learning_list import render_learning_list
from src.cards.models import CardEntry
from src.config import Cards


@pytest.mark.parametrize("language", ["en", "es"])
def test_learning_list_matches_printed_card(tmp_path, language):
    card = CardEntry(
        lemma="watch",
        pos="NOUN, VERB",
        rank=1,
        grammatical_forms="watch / watches / watched",
        learning_form="watch",
        ipa="wɒtʃ",
        translation_ru="смотреть, часы",
        translation_en="watch",
        observed_forms={"watch": 5, "watches": 3, "watched": 1},
    )
    config = Cards()
    docx = tmp_path / "cards.docx"
    xlsx = tmp_path / "list.xlsx"
    render([card], Path(__file__).parents[1] / "table.docx", docx, config, language)
    render_learning_list([card], xlsx, config, language)
    document = Document(docx)
    front = document.tables[0].cell(0, 0).text
    back = [p.text for p in document.tables[1].cell(7, 0).paragraphs]
    book = load_workbook(xlsx)
    sheet = book.active
    assert sheet.max_column == 6
    assert [c.value for c in sheet[1]][4:] == ["Лемма", "Знаю"]
    assert sheet.max_row == 2
    row = [c.value for c in sheet[2]]
    assert row[0] == front
    assert row[1] == back[0]
    translations = 2 if language == "es" else 1
    assert row[2] == "\n".join(back[1 : 1 + translations])
    assert (row[3] or "") == (back[-1] if len(back) > 1 + translations else "")
    assert row[4] == card.lemma and not row[5]
    assert sheet.freeze_panes == "A2"
    book.close()
    # openpyxl accepts metadata that Windows/Excel rejects; inspect the package too.
    for artifact in (xlsx, docx):
        with ZipFile(artifact) as archive:
            core = ElementTree.fromstring(archive.read("docProps/core.xml"))
        assert all("{http://www.w3.org/XML/1998/namespace}space" not in node.attrib for node in core.iter())


def test_learning_list_text_cannot_become_excel_formula(tmp_path):
    card = CardEntry("=1+1", "NOUN", 1, "", "=1+1", "-", "=2+2", "", {})
    target = tmp_path / "list.xlsx"
    render_learning_list([card], target, Cards(), "en")
    book = load_workbook(target)
    assert book.active["A2"].data_type == "s"
    assert book.active["C2"].data_type == "s"
    book.close()
