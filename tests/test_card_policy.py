from openpyxl import Workbook, load_workbook

from src.cards.known_words import append_words, import_marked_words, import_word_file, read_known_words
from src.cards.level import frequent_lemmas
from src.cards.selector import CardPolicy, forecast, select_cards
from src.config import TranslationSettings


def lemma(id_, word, pos, count):
    return {
        "id": id_, "lemma": word, "pos": pos, "count": count, "observed_forms": {word: count},
        "contexts": [], "morph_variants": [], "chunk_count": 1,
    }


def dataset():
    lemmas = [
        lemma("1", "και", "CCONJ", 50),
        lemma("2", "δρόμος", "NOUN", 30),
        lemma("3", "νίκος", "PROPN", 20),
        lemma("4", "league", "X", 15),
        lemma("5", "δύο", "NUM", 12),
        lemma("6", "γράφω", "VERB", 10),
        lemma("7", "σπίτι", "NOUN", 8),
    ]
    forms = [{"id": f"f{r['id']}", "form": r["lemma"], "lemma": r["lemma"], "pos": r["pos"], "count": r["count"],
              "reference_frequency": 0.001} for r in lemmas]
    return {"lemmas": lemmas, "forms": forms, "lemma_grammar": [], "ipa": [], "translations": []}


def words(cards):
    return [card.lemma for card in cards]


def test_filters_skip_categories_and_other_alphabets():
    policy = CardPolicy(True, True, True, True)
    assert words(select_cards(dataset(), 100, 0, 1, "el", policy=policy)[0]) == ["δρόμος", "γράφω", "σπίτι"]
    assert len(select_cards(dataset(), 100, 0, 1, "el")[0]) == 7


def test_assumed_vocabulary_and_card_limit():
    policy = CardPolicy(True, True, True, True, max_cards=1, assumed_known=frozenset({"δρόμος"}))
    assert words(select_cards(dataset(), 100, 0, 1, "el", policy=policy)[0]) == ["γράφω"]


def test_forecast_ignores_the_limit():
    policy = CardPolicy(max_cards=1)
    settings = TranslationSettings(specificity_threshold=0, min_book_occurrences=1)
    result = forecast(dataset(), settings, "el", set(), policy)
    assert result[95] == 7 and result[70] < result[95]


class Lexicon:
    def lemmas_of(self, form):
        return {"είπε": {"λέω"}, "τησ": set(), "της": {"ο"}, "δρόμου": {"δρόμος"}}.get(form, set())


def test_frequent_lemmas_restore_final_sigma_and_merge_forms():
    lookup = lambda language, n: ["και", "της", "είπε", "λέει", "123", "δρόμου"]
    assert frequent_lemmas("el", 2, Lexicon(), lookup) == {"και", "ο"}
    assert frequent_lemmas("el", 0, Lexicon(), lookup) == set()
    assert "λέω" in frequent_lemmas("el", 3, Lexicon(), lookup)


def test_import_words_from_text_file_and_study_list(tmp_path):
    dictionary = tmp_path / "my_dictionary_el.xlsx"
    append_words(dictionary, ["ο"])
    source = tmp_path / "words.txt"
    source.write_text("Δρόμος\nο\n\nγράφω\n", encoding="utf-8")
    assert import_word_file(source, dictionary) == 2
    assert read_known_words(dictionary) == {"known words", "ο", "δρόμος", "γράφω"}

    study_list = tmp_path / "greek_list.xlsx"
    book = Workbook()
    book.active.append(["Word", "Transcription", "Translation", "Word forms", "Lemma", "Known"])
    book.active.append(["το σπίτι", "", "", "", "σπίτι", "x"])
    book.active.append(["η πόλη", "", "", "", "πόλη", None])
    book.active.append(["ο δρόμος", "", "", "", "δρόμος", "+"])
    book.save(study_list)
    assert import_marked_words(study_list, dictionary) == 1
    sheet = load_workbook(dictionary).active
    assert [row[0] for row in sheet.iter_rows(values_only=True)][-1] == "σπίτι"
